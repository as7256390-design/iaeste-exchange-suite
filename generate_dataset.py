import openpyxl, json, re

# Normalize country name
def norm_name(name):
    if not name:
        return ''
    n = str(name).strip()
    # Strip suffixes like ', KU', ', LC', etc for country grouping if needed, or keep canonical IAESTE name
    return n

# 1. 2019 data
wb19 = openpyxl.load_workbook('Activity Report 2019/World Exchange 2019.xlsx', data_only=True)
ws19 = wb19['Combined SentReceived']
data_2019 = {}
for r in list(ws19.iter_rows(values_only=True))[2:]:
    name = r[0]
    if not name or 'total' in str(name).lower():
        continue
    sent = r[5]
    recv = r[6]
    sent_val = float(sent) if sent not in ('-', None, '#N/A') else 0.0
    recv_val = float(recv) if recv not in ('-', None, '#N/A') else 0.0
    data_2019[name.strip()] = {'outgoing': sent_val, 'incoming': recv_val}

# 2. 2020 - 2024 Trend Data
wb_trend = openpyxl.load_workbook('Activity Report 2024/trend_ar_2020___2024_2024-12-05T15_27_20.812072Z.xlsx', data_only=True)
ws_trend = wb_trend['data']
trend_by_year = {}
for r in list(ws_trend.iter_rows(values_only=True))[1:]:
    year = str(r[0]).strip()
    comm = str(r[1]).strip()
    if not year or not comm:
        continue
    outg = float(r[2]) if r[2] is not None else 0.0
    inc = float(r[3]) if r[3] is not None else 0.0
    if year not in trend_by_year:
        trend_by_year[year] = {}
    trend_by_year[year][comm] = {'outgoing': outg, 'incoming': inc}

# 3. 2024 Offer Totals
wb_ot = openpyxl.load_workbook('Activity Report 2024/offer_totals_2024-12-05T15_24_41.922352Z.xlsx', data_only=True)
ws_ot = wb_ot.active
ot_2024 = {}
for r in list(ws_ot.iter_rows(values_only=True))[1:]:
    comm = str(r[0]).strip()
    ot_2024[comm] = {
        'published': float(r[1] or 0),
        'nominated': float(r[2] or 0),
        'accepted': float(r[3] or 0),
        'unused': float(r[4] or 0),
        'cancelled': float(r[5] or 0),
        'onsite': float(r[9] or 0),
        'remote': float(r[10] or 0),
        'hybrid': float(r[11] or 0),
    }

# 4. Bilateral Exchange with India
wb_ex = openpyxl.load_workbook('Activity Report 2024/exchange_sent___committee_2024-12-05T15_25_15.254135Z.xlsx', data_only=True)
ws_ex = wb_ex.active
india_sent_to = {}
india_recv_from = {}
for r in list(ws_ex.iter_rows(values_only=True))[1:]:
    sender = str(r[0]).strip()
    receiver = str(r[1]).strip()
    count = float(r[2] or 0)
    if sender.lower().startswith('india'):
        india_sent_to[receiver] = count
    if receiver.lower().startswith('india'):
        india_recv_from[sender] = count

# 5. Disciplines
wb_s = openpyxl.load_workbook('Activity Report 2024/general_discipline___sent_2024-12-05T15_26_48.225014Z.xlsx', data_only=True)
ws_s = wb_s['data']
disc_sent = {}
for r in list(ws_s.iter_rows(values_only=True))[1:]:
    comm = str(r[0]).strip()
    disc = str(r[1]).strip()
    cnt = float(r[2] or 0)
    if comm not in disc_sent:
        disc_sent[comm] = {}
    disc_sent[comm][disc] = disc_sent[comm].get(disc, 0) + cnt

wb_r = openpyxl.load_workbook('Activity Report 2024/general_discipline___received_2024-12-05T15_26_57.994772Z.xlsx', data_only=True)
ws_r = wb_r['data']
disc_recv = {}
for r in list(ws_r.iter_rows(values_only=True))[1:]:
    comm = str(r[0]).strip()
    disc = str(r[1]).strip()
    cnt = float(r[2] or 0)
    if comm not in disc_recv:
        disc_recv[comm] = {}
    disc_recv[comm][disc] = disc_recv[comm].get(disc, 0) + cnt

print("Data parsed successfully.")
